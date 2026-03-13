import json
import math
import os
import sys
import csv
from datetime import datetime

import threading
import time
import logging
from collections import deque
from typing import List, Dict, Any, Optional, Tuple, Union
import serial
import serial.tools.list_ports
import cv2
import numpy as np

from backend.core import config
import openpyxl
from openpyxl.styles import PatternFill
from backend.services import db_service

# Configure logging
logger = logging.getLogger("itms.sensors")

# ------------------------------------------------------------------------------
# GLOBAL STATE
# ------------------------------------------------------------------------------

# Configuration defaults
DEFAULT_PORT = "/dev/cu.usbserial-0001"
DEFAULT_BAUD = 115200

# Known USB-Serial chip identifiers for auto-detection
_KNOWN_SERIAL_KEYWORDS = [
    "CH340", "CH341",        # Common cheap USB-UART chips
    "CP210", "CP2102",       # Silicon Labs
    "FTDI", "FT232",         # FTDI chips
    "USB-SERIAL", "USB Serial",
    "Arduino",
    "ESP32", "ESP",
    "usbserial", "ttyUSB", "ttyACM",
    "USB2.0-Ser",            # Prolific PL2303
    "PL2303",
    "usbmodem",
]
METERS_PER_STEP_DEFAULT = 0.004

# Data Buffers (for history/plots)
BUFFER_SIZE = 300
_packet_count = 0 # Diagnostic counter
_history_lock = threading.Lock()
_history_buffer: deque = deque(maxlen=BUFFER_SIZE)

# SLAM / Map Data
_slam_lock = threading.Lock()
_map_points: List[Tuple[float, float]] = []  # List of (x, y) tuples
_robot_y: float = 0.0
_encoder_steps: int = 0
_last_encoder_steps: int = 0
_meters_per_step: float = METERS_PER_STEP_DEFAULT
_reset_next_sample: bool = False

# Dead Reckoning / Hybrid GPS+Encoder
_gps_origin: Optional[Tuple[float, float]] = None   # (lat, lon) captured from first valid GPS fix
_heading_deg: float = 0.0                            # Track heading in degrees (0 = North, 90 = East)

# Thresholds & Alerts
_thresholds_lock = threading.Lock()
_thresholds: Dict[str, float] = {
    "acc_x": 2.0, "acc_y": 2.0, "acc_z": 2.0,
    "gyro_x": 5.0, "gyro_y": 5.0, "gyro_z": 5.0,
    "twist": 10.0,
    "clearance": 1.5, # SLAM clearance
    "uml": 0.15, # Urgent Maintenance Limit (Red)
    "pml": 0.4, # Planned Maintenance Limit (Yellow/Orange)
}
_alerts: deque = deque(maxlen=50)

# Infringement Logging
_infringement_file = None
_infringement_filename = None

# Current Reading (snapshot)
_current_data: Dict[str, Any] = {
    "accel": {"x": 0, "y": 0, "z": 0},
    "gyro": {"x": 0, "y": 0, "z": 0},
    "gps": {"lat": 0, "lon": 0},
    "geometry": {"gauge": 1676, "twist": 0, "crosslevel": 0},
    "encoder_steps": 0,
    "lidar_mm": -1,
    "timestamp": 0
}

# Control Flags
_running = False
_recording = False
_serial_thread: Optional[threading.Thread] = None

# Serial Port
_ser: Optional[serial.Serial] = None

# Recording Handles
_log_file = None
_video_writer1 = None
_video_writer2 = None
_writer1_shape = None
_writer2_shape = None
_record_start_time = 0
_current_session_id: Optional[int] = None

# Callbacks for WebSocket broadcasting
_broadcast_callback = None

def set_broadcast_callback(cb):
    """Set a callback function(data_dict) to push data to WebSockets."""
    global _broadcast_callback
    _broadcast_callback = cb

# ------------------------------------------------------------------------------
# CORE SERIAL LOGIC
# ------------------------------------------------------------------------------

def auto_detect_serial_port() -> Optional[str]:
    """
    Scan all available serial ports and return the device path of the first
    port whose description or hardware ID matches a known USB-Serial chip.
    Works cross-platform: COMx (Windows), /dev/cu.* (macOS), /dev/ttyUSB* (Linux).
    Returns None if no suitable port is found.
    """
    ports = list(serial.tools.list_ports.comports())

    if not ports:
        logger.warning("Auto-detect: No serial ports found on this system.")
        return None

    # Log every port for diagnostics
    logger.info(f"Auto-detect: Found {len(ports)} serial port(s):")
    for p in ports:
        logger.info(f"  • {p.device}  desc='{p.description}'  hwid='{p.hwid}'")

    # First pass – match against known USB-serial keywords
    for p in ports:
        searchable = f"{p.description} {p.hwid}".upper()
        for keyword in _KNOWN_SERIAL_KEYWORDS:
            if keyword.upper() in searchable:
                logger.info(f"Auto-detect: Matched '{keyword}' → {p.device}")
                return p.device

    # Second pass – on Linux/macOS pick the first /dev/ttyUSB* or /dev/cu.usb*
    for p in ports:
        dev = p.device.lower()
        if "ttyusb" in dev or "ttyacm" in dev or "cu.usb" in dev or "cu.wch" in dev:
            logger.info(f"Auto-detect: Falling back to USB-style device → {p.device}")
            return p.device

    # Third pass – pick first non-Bluetooth port on Windows (COMx)
    for p in ports:
        desc_upper = p.description.upper()
        if "BLUETOOTH" not in desc_upper and p.device.upper().startswith("COM"):
            logger.info(f"Auto-detect: Falling back to first non-BT COM port → {p.device}")
            return p.device

    logger.warning("Auto-detect: Could not identify a suitable serial device.")
    return None


def start_acquisition(port: str = DEFAULT_PORT, baud: int = DEFAULT_BAUD, meters_per_step: float = METERS_PER_STEP_DEFAULT):
    global _running, _serial_thread, _meters_per_step
    
    if _running:
        logger.warning("Acquisition already running.")
        return

    # ── Resolve serial port ──────────────────────────────────────────
    resolved_port = port
    if not port or port.upper() == "AUTO":
        logger.info("SERIAL_PORT not set or set to AUTO – running auto-detection…")
        detected = auto_detect_serial_port()
        if detected:
            resolved_port = detected
            logger.info(f"Using auto-detected serial port: {resolved_port}")
        else:
            logger.error("No serial device detected. Please connect a device or set SERIAL_PORT in .env")
            return
    else:
        logger.info(f"Using manually configured serial port: {resolved_port}")

    _meters_per_step = meters_per_step
    _running = True

    _serial_thread = threading.Thread(target=_serial_loop, args=(resolved_port, baud), daemon=True)
    _serial_thread.start()
    logger.info(f"Started acquisition on {resolved_port} @ {baud}")

def stop_acquisition():
    global _running
    _running = False
    if _serial_thread:
        _serial_thread.join(timeout=2.0)
    
    _close_serial()
    stop_db_session() # Ensure session ends if acquisition stops
    logger.info("Stopped acquisition.")

def start_db_session(data_source: str):
    """Start a persistent database session."""
    global _current_session_id
    if _current_session_id is None:
        _current_session_id = db_service.db.start_session(data_source)
        logger.info(f"Database session started: {_current_session_id}")

def stop_db_session():
    """End the current persistent database session."""
    global _current_session_id, _robot_y
    if _current_session_id is not None:
        db_service.db.end_session(_current_session_id, _robot_y)
        logger.info(f"Database session stopped: {_current_session_id}")
        _current_session_id = None

# def _mock_loop():
#     """
#     Simulation loop that generates random sensor data.
#     Used for testing UI behavior and safety gating when hardware is not connected.
#     """
#     global _running
#     import random
#     
#     steps = 0
#     while _running:
#         steps += 10
#         data = {
#             "accel": {
#                 "x": random.uniform(-0.5, 0.5),
#                 "y": random.uniform(-0.5, 0.5),
#                 "z": 9.8 + random.uniform(-0.2, 0.2)
#             },
#             "gyro": {
#                 "x": random.uniform(-1, 1),
#                 "y": random.uniform(-1, 1),
#                 "z": random.uniform(-1, 1)
#             },
#             "encoder_steps": steps,
#             "lidar1_mm": 1000 + random.randint(-50, 50),
#             "lidar2_mm": 1000 + random.randint(-50, 50),
#             "lidar3_mm": 2500 + random.randint(-20, 20),
#             "y": steps * 0.004
#         }
#         _process_data(data)
#         time.sleep(0.1)

def _close_serial():
    global _ser
    if _ser:
        try:
            _ser.close()
        except Exception as e:
            logger.error(f"Error closing serial: {e}")
    _ser = None

def reset_system():
    """Reset encoder distance, map, history, and dead reckoning origin."""
    global _robot_y, _encoder_steps, _last_encoder_steps, _map_points, _history_buffer, _gps_origin
    
    # Send 'R' to microcontroller if supported
    if _ser:
        try:
            _ser.write(b"R")
        except Exception:
            pass

    with _slam_lock:
        _robot_y = 0.0
        global _reset_next_sample
        _reset_next_sample = True
        _map_points.clear()
        _gps_origin = None  # Force re-capture on next valid GPS packet
    
    with _history_lock:
        _history_buffer.clear()
        
    logger.info("System reset performed (dead reckoning origin cleared).")

def _serial_loop(port, baud):
    global _ser, _running, _current_data, _robot_y, _last_encoder_steps
    
    # Attempt connection
    try:
        _ser = serial.Serial(port, baud, timeout=1)
        logger.info(f"Connected to {port}")
    except Exception as e:
        logger.error(f"Failed to connect to {port}: {e}")
        _running = False
        return

    while _running:
        try:
            if not _ser:
                break
                
            line = _ser.readline().decode("utf-8", errors="ignore").strip()
            if not line:
                continue

            # Log raw if recording -> CHANGED: Now logging structured CSV in _process_data
            # if _recording:
            #     _log_raw_serial(line)
            
            # Use raw line for infringement logging if needed, or structured data below

            # Parse JSON
            if line.startswith("{") and line.endswith("}"):
                try:
                    data = json.loads(line)
                    # Debug print to confirm data flow
                   # print(f"DEBUG SENSORS: Valid JSON received: {str(data)[:60]}...") 
                    _process_data(data)
                except json.JSONDecodeError:
                   # print(f"DEBUG SENSORS: JSON Decode Error on line: {line}")
                    pass
        except Exception as e:
            logger.error(f"Serial loop error: {e}")
            time.sleep(0.1)

    _close_serial()



def _classify_distance(dist_m: float, uml: float, pml: float) -> str:
    """
    IF distance == 0: classification = "CBML" (Green)
    ELIF distance < UML_threshold: classification = "UML" (Red)
    ELIF distance < PML_threshold: classification = "PML" (Yellow)
    ELSE: classification = "CBML" (Green)
    """
    # Treat 0 or invalid as safe (CBML) based on prompt
    if dist_m <= 0.001: 
        return "CBML"
    if dist_m < uml:
        return "UML"
    if dist_m < pml:
        return "PML"
    return "CBML"

def _process_data(data: Dict[str, Any]):
    global _current_data, _robot_y, _last_encoder_steps, _encoder_steps, _infringement_file, _reset_next_sample, _packet_count

    timestamp = time.time()
    data["timestamp"] = timestamp
    
    # Update global snapshot
    _current_data = data

    # 1. Update SLAM / Odometry
    steps = data.get("encoder_steps", 0)
    # Lidar readings (mm)
    l1_mm = data.get("lidar1_mm", 0)
    l2_mm = data.get("lidar2_mm", 0)
    l3_mm = data.get("lidar3_mm", 0)
    
    # Also support legacy single lidar if provided
    lidar_mm = data.get("lidar_mm", -1)
    


    with _slam_lock:
        if _reset_next_sample:
            # Sync to current hardware state
            _robot_y = 0.0
            _last_encoder_steps = steps
            _encoder_steps = steps
            _reset_next_sample = False
            return # Skip this sample processing to avoid jump

    delta = steps - _last_encoder_steps
    dist_delta = delta * _meters_per_step
    
    # 2. Derive Geometry & Check Thresholds
    ax = data.get("accel", {}).get("x", 0)
    ay = data.get("accel", {}).get("y", 0)
    gx = data.get("gyro", {}).get("x", 0)
    
    twist_val = abs(gx) 
    crosslevel_val = abs(ax)
    
    data["geometry"] = {
        "gauge": 1676,
        "twist": twist_val,
        "crosslevel": crosslevel_val
    }
    
    _check_thresholds(data, _robot_y)
    
    with _slam_lock:
        _robot_y += dist_delta  # pyre-ignore[58]
        _last_encoder_steps = steps
        _encoder_steps = steps
        
        # Add to map (using lidar1/left as primary for SLAM map if available, else legacy)
        # Using lidar1 (Left) -> Negative X
        valid_map_point = False
        if l1_mm > 50 and l1_mm < 4000:
             _map_points.append((-l1_mm / 1000.0, _robot_y))
             valid_map_point = True
        elif lidar_mm > 50 and lidar_mm < 4000: # Fallback
             _map_points.append((lidar_mm / 1000.0, _robot_y))

    # Convert to meters
    l1_m = l1_mm / 1000.0
    l2_m = l2_mm / 1000.0
    l3_m = l3_mm / 1000.0

    # classification
    with _thresholds_lock:
        uml = _thresholds.get("uml", 1.0)
        pml = _thresholds.get("pml", 0.8)
    
    c1 = _classify_distance(l1_m, uml, pml)
    c2 = _classify_distance(l2_m, uml, pml)
    c3 = _classify_distance(l3_m, uml, pml)
    
    # Log to CSV if recording
    if _infringement_file:
        try:
            # Format timestamp as HH:MM:SS
            ts_str = time.strftime("%H:%M:%S", time.localtime(timestamp))
            line = f"{ts_str},{_robot_y:.4f},{l1_m:.4f},{l2_m:.4f},{l3_m:.4f},{c1},{c2},{c3}\n"
            _infringement_file.write(line)
            _infringement_file.flush()
        except:
            pass

    # Log Acceleration/IMU Data to CSV if recording
    if _log_file:
        try:
            ts_str2 = time.strftime("%H:%M:%S", time.localtime(timestamp))
            # Format: Timestamp,Chainage,AX,AY,AZ,GX,GY,GZ
            # Gyro might be in data['gyro']
            gx = data.get("gyro", {}).get("x", 0)
            gy = data.get("gyro", {}).get("y", 0)
            gz = data.get("gyro", {}).get("z", 0)
            
            row = f"{ts_str2},{_robot_y:.4f},{ax:.4f},{ay:.4f},{data.get('accel', {}).get('z', 0):.4f},{gx:.4f},{gy:.4f},{gz:.4f}\n"
            _log_file.write(row)
            _log_file.flush()
        except Exception as e:
            pass

    # ── Dead Reckoning: compute estimated GPS from encoder displacement ──
    global _gps_origin
    raw_gps = data.get("gps", {})
    raw_lat = raw_gps.get("lat", 0)
    raw_lon = raw_gps.get("lon", 0)

    # Capture GPS origin on first valid fix
    if _gps_origin is None and raw_lat != 0 and raw_lon != 0:
        _gps_origin = (raw_lat, raw_lon)
        logger.info(f"Dead Reckoning: GPS origin captured at ({raw_lat:.6f}, {raw_lon:.6f})")

    # Compute dead-reckoned position
    dr_gps = {}
    if _gps_origin is not None:
        # 1 degree latitude  ≈ 111,320 meters
        # 1 degree longitude ≈ 111,320 * cos(lat) meters
        lat_offset = (_robot_y * math.cos(math.radians(_heading_deg))) / 111320.0
        lon_offset = (_robot_y * math.sin(math.radians(_heading_deg))) / (111320.0 * math.cos(math.radians(_gps_origin[0])))
        dr_gps = {
            "lat": _gps_origin[0] + lat_offset,
            "lon": _gps_origin[1] + lon_offset
        }

    # 3. Add to History Buffer
    processed_point = {
        "timestamp": timestamp,
        "accel": data.get("accel", {}),
        "gyro": data.get("gyro", {}),
        "gps": raw_gps,
        "dr_gps": dr_gps,
        "geometry": data.get("geometry", {}),
        "y": _robot_y,
        "lidar_m": l1_m if l1_m > 0 else (lidar_mm / 1000.0 if lidar_mm > 0 else 0),
        "lidar": {
            "l1": l1_m,
            "l2": l2_m,
            "l3": l3_m
        },
        "classification": {
            "l1": c1, "l2": c2, "l3": c3
        }
    }
    
    with _history_lock:
        _history_buffer.append(processed_point)

    # 3.5 Persistent Database Logging
    global _current_session_id
    if _current_session_id is not None:
        try:
            db_service.db.log_telemetry(_current_session_id, processed_point)
        except Exception as e:
            logger.error(f"Failed to log telemetry to DB: {e}")

    # 3. Broadcast to WebSockets
    global _packet_count
    _packet_count += 1  # pyre-ignore[58]
    if _packet_count % 50 == 0:
        logger.info(f"Telemetry Heartbeat: Broadcasted {_packet_count} packets. Current Y: {processed_point['y']:.3f}m")

    if _broadcast_callback:
        _broadcast_callback({
            "type": "update",
            "data": processed_point
        })

def _check_thresholds(data: Dict[str, Any], dist_y: float):
    # Check IMU
    acc = data.get("accel", {})
    gyro = data.get("gyro", {})
    
    with _thresholds_lock:
        th = _thresholds.copy()
        
    for axis in ["x", "y", "z"]:
        val = abs(acc.get(axis, 0))
        limit = th.get(f"acc_{axis}", 2.0)
        if val > limit:
            _add_alert("Acceleration", f"Acc {axis.upper()} {val:.2f} > {limit}", dist_y)
            
        val_g = abs(gyro.get(axis, 0))
        limit_g = th.get(f"gyro_{axis}", 5.0)
        if val_g > limit_g:
            _add_alert("Gyroscope", f"Gyro {axis.upper()} {val_g:.2f} > {limit_g}", dist_y)

def _add_alert(category: str, msg: str, dist_y: float):
    alert = {
        "time": time.strftime("%H:%M:%S"),
        "y": dist_y,
        "category": category,
        "message": msg
    }
    with _history_lock:
        _alerts.append(alert)
        # Notify via broadcast if needed, specific alert type
        if _broadcast_callback:
            _broadcast_callback({"type": "alert", "data": alert})

# ------------------------------------------------------------------------------
# RECORDING LOGIC
# ------------------------------------------------------------------------------

# ------------------------------------------------------------------------------
# CAMERA HELPERS (invoked by server.py generators)
# ------------------------------------------------------------------------------

def _ensure_directories():
    root = config.STORAGE_DIR
    subdirs = [
        "Acceleration", 
        "video_recording/RearWindow", 
        "video_recording/RailCondition", 
        "video_recording/TrackGeometry", 
        "video_recording/RailProfile",
        "video_recording/ConditionMonitoring",
        "Recordings", 
        "Infringements"
    ]
    for d in subdirs:
        path = os.path.join(root, d)
        os.makedirs(path, exist_ok=True)
    return root

def record_frame(cam_index, frame):
    """
    Called by the video stream generator to save frames if recording is active.
    Initializing writers on the fly to handle dynamic resolution.
    """
    global _video_writer1, _video_writer2, _recording
    global _writer1_shape, _writer2_shape

    if not _recording:
        return

    # Helper to init writer
    def init_writer(outfile):
        h, w = frame.shape[:2]
        fourcc = cv2.VideoWriter_fourcc(*'mp4v')
        return cv2.VideoWriter(outfile, fourcc, 30.0, (w, h)), (w, h)

    root = _ensure_directories()
    ts = time.strftime("%Y%m%d_%H%M%S")

    # Lazy init writers
    if cam_index == 0: # RearWindow
        if _video_writer1 is None:
            path = os.path.join(root, "video_recording", "RearWindow", f"cam0_rear_{ts}.mp4")
            _video_writer1, _writer1_shape = init_writer(path)
            logger.info(f"Recording Cam 0 to {path}")
        
        # Resize if camera resolution changed mid-stream
        if frame.shape[:2] != (_writer1_shape[1], _writer1_shape[0]):
            frame = cv2.resize(frame, _writer1_shape)
            
        if _video_writer1 is not None:
            _video_writer1.write(frame)
        
    elif cam_index == 1: # RailCondition / Profile
        if _video_writer2 is None:
            path = os.path.join(root, "video_recording", "RailCondition", f"cam1_{ts}.mp4")
            _video_writer2, _writer2_shape = init_writer(path)
            logger.info(f"Recording Cam 1 to {path}")
            
        if frame.shape[:2] != (_writer2_shape[1], _writer2_shape[0]):
            frame = cv2.resize(frame, _writer2_shape)
            
        if _video_writer2 is not None:
            _video_writer2.write(frame)
    
    # Example logic for other cameras if needed
    # else:
    #    ...

# Also update start_recording to use the data folder for serial logs
def start_recording(file_prefix="session"):
    global _recording, _log_file, _record_start_time, _video_writer1, _video_writer2
    if _recording:
        return

    ts_str = time.strftime("%Y%m%d_%H%M%S")
    _record_start_time = time.time()
    
    root = _ensure_directories()
    
    # 1. Open Serial Log (Acceleration Data) -> NOW CSV
    filename = os.path.join(root, "Acceleration", f"{file_prefix}_{ts_str}_serial.csv")
    try:
        _log_file = open(filename, "w")
        # Write CSV Header
        _log_file.write("Timestamp,Chainage,AccX,AccY,AccZ,GyroX,GyroY,GyroZ\n")
    except Exception as e:
        logger.error(f"Failed to open log file: {e}")

    # Reset video writers so they re-init with new files
    global _writer1_shape, _writer2_shape
    _video_writer1 = None
    _video_writer2 = None
    _writer1_shape = None
    _writer2_shape = None
    
    _recording = True
    logger.info(f"Recording started. Log: {filename}")
    logger.info(f"Recording started. Log: {filename}")

    # Start Infringement CSV
    _infringement_filename = os.path.join(root, "Infringements", f"{file_prefix}_{ts_str}_infringements.csv")
    try:
        global _infringement_file
        _infringement_file = open(_infringement_filename, "w")
        # Header
        _infringement_file.write("timestamp,y,lidar1_m,lidar2_m,lidar3_m,class_1,class_2,class_3\n")
    except Exception as e:
        logger.error(f"Failed to open infringement log: {e}")

    return filename

def stop_recording():
    global _recording, _log_file, _video_writer1, _video_writer2
    global _writer1_shape, _writer2_shape
    _recording = False
    
    if _log_file:
        _log_file.close()
        _log_file = None
        
    if _video_writer1:
        _video_writer1.release()
        _video_writer1 = None
        _writer1_shape = None
        
    if _video_writer2:
        _video_writer2.release()
        _video_writer2 = None
        _writer2_shape = None

    global _infringement_file
    if _infringement_file:
        _infringement_file.close()
        _infringement_file = None

    logger.info("Recording stopped.")

def _log_raw_serial(line):
    if _log_file:
        t = time.time() - _record_start_time
        _log_file.write(f"[{t:.3f}] {line}\n")
        _log_file.flush()

# ------------------------------------------------------------------------------
# GETTERS
# ------------------------------------------------------------------------------

def is_running():
    return _running

def is_recording():
    return _recording

def get_latest_state():
    with _slam_lock:
        return {
            "running": _running,
            "recording": _recording,
            "y": _robot_y,
            "encoder_steps": _encoder_steps,
            "last_packet": _current_data
        }

def get_history():
    with _history_lock:
        return list(_history_buffer)

def get_map_points():
    with _slam_lock:
        return list(_map_points)

def get_alerts():
    with _history_lock:
        return list(_alerts)

def set_thresholds(new_th: Dict[str, float]):
    with _thresholds_lock:
        _thresholds.update(new_th)
    logger.info(f"Thresholds updated: {new_th}")

def get_thresholds():
    with _thresholds_lock:
        return _thresholds.copy()

def export_infringement_history():
    """Export current memory buffer to Excel with color coding."""
    # root = os.path.join(os.getcwd(), "data", "Infringements")
    root = os.path.join(config.STORAGE_DIR, "Infringements")
    os.makedirs(root, exist_ok=True)
    ts_str = time.strftime("%Y%m%d_%H%M%S")
    filename = os.path.join(root, f"export_{ts_str}.xlsx")
    
    with _thresholds_lock:
        uml = _thresholds.get("uml", 1.0)
        pml = _thresholds.get("pml", 0.8)

    with _history_lock:
        data_to_write = list(_history_buffer)

    # Create Workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Infringements"

    # Define Fills
    fill_red = PatternFill(start_color="FF9999", end_color="FF9999", fill_type="solid")
    fill_orange = PatternFill(start_color="FFCC99", end_color="FFCC99", fill_type="solid")
    fill_green = PatternFill(start_color="99FF99", end_color="99FF99", fill_type="solid")

    def get_fill(cls_str):
        if cls_str == "UML": return fill_red
        if cls_str == "PML": return fill_orange
        return fill_green # CBML

    # Header
    headers = ["Timestamp", "Chainage (y)", "Lidar1 (m)", "Lidar2 (m)", "Lidar3 (m)", "Classification"]
    ws.append(headers)

    for item in data_to_write:
        ts = item.get("timestamp", 0)
        y = item.get("y", 0)
        
        l_obj = item.get("lidar", {})
        l1 = l_obj.get("l1", 0)
        l2 = l_obj.get("l2", 0)
        l3 = l_obj.get("l3", 0)
        
        # Calculate individual classes
        c1 = _classify_distance(l1, uml, pml)
        c2 = _classify_distance(l2, uml, pml)
        c3 = _classify_distance(l3, uml, pml)
        
        # Aggregation Logic
        # Priority: UML > PML > CBML
        # If any is UML -> UML
        # Else if any is PML -> PML
        # Else -> CBML
        
        final_class = "CBML"
        if "UML" in [c1, c2, c3]:
            final_class = "UML"
        elif "PML" in [c1, c2, c3]:
            final_class = "PML"
        
        ts_str = time.strftime("%H:%M:%S", time.localtime(ts))
        
        # Row Data
        row = [ts_str, y, l1, l2, l3, final_class]
        ws.append(row)

        # Apply Colors to the last column (Classification)
        curr_row = ws.max_row
        ws.cell(row=curr_row, column=6).fill = get_fill(final_class)

    wb.save(filename)
    logger.info(f"Exported colored Excel to {filename}")
    return filename
