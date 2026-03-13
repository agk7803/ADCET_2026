import os
import time
import json
import cv2
import logging
from typing import Optional, List, Dict, Any
from fastapi import APIRouter, WebSocket, WebSocketDisconnect, HTTPException, Body, Query
from fastapi.responses import StreamingResponse
from pydantic import BaseModel


from backend import sensors
from fastapi.responses import FileResponse
from backend import simulator
from backend.core import config
from backend.services import sensor_service, cv_service, db_service, report_service
from backend.services import report_session

# Global State from original server.py
CONDITION_MODE = "laser"

router = APIRouter()
logger = logging.getLogger("itms.api")

def get_current_chainage():
    return sensors.get_latest_state().get('y', 0.0)

# Models for API
class AccelPoint(BaseModel):
    time: str
    rawTime: float
    ax: float
    ay: float
    az: float
    chainage: float
    status: Optional[str] = "CBML"

class AccelExportRequest(BaseModel):
    data: List[AccelPoint]

class ModeUpdate(BaseModel):
    mode: str 

class ThresholdUpdate(BaseModel):
    threshold: int

class ModelSelect(BaseModel):
    path: str

class CameraControl(BaseModel):
    index: int

class SensorVector(BaseModel):
    x: float
    y: float
    z: float

class SensorGPS(BaseModel):
    lat: float
    lon: float

class SensorPacket(BaseModel):
    accel: SensorVector
    gyro: SensorVector
    gps: SensorGPS
    encoder_steps: int
    lidar1_mm: int
    lidar2_mm: int
    lidar3_mm: int

# Endpoints
@router.api_route("/", methods=["GET", "HEAD"])
def root():
    return {"status": "ITMS Backend Online", "running": sensors.is_running()}

@router.get("/health")
async def health_check():
    return {
        "status": "ok",
        "sensors_running": sensors.is_running(),
        "data_source": config.DATA_SOURCE,
        "timestamp": time.time()
    }

@router.post("/connect")
def connect_system(
    port: str = Body(None, embed=True), # Optional override
    baud: int = Body(None, embed=True), # Optional override
    meters_per_step: float = Body(0.004, embed=True)
):
    try:
        # Use provided overrides or fallback to config defaults
        active_port = port if port else config.SERIAL_PORT
        active_baud = baud if baud else config.SERIAL_BAUD
        
        if config.DATA_SOURCE == "SIMULATOR":
            if not simulator.is_running():
                # Default simulator to 20Hz as before
                simulator.start_simulator(rate_hz=20)
                sensors.start_db_session("SIMULATOR")
            report_session.start_report_session()
            return {"status": "connected", "mode": "SIMULATOR"}
            
        elif config.DATA_SOURCE == "SERIAL":
            if simulator.is_running():
                return {"status": "connected", "source": "simulator"}
            sensors.start_acquisition(active_port, active_baud, meters_per_step)
            sensors.start_db_session("SERIAL")
            report_session.start_report_session()
            return {"status": "connected", "mode": "SERIAL", "config": {"port": active_port, "baud": active_baud}}
            
        elif config.DATA_SOURCE == "WIFI":
            # Just open the gate for incoming HTTP POST traffic
            sensors._running = True
            sensors.start_db_session("WIFI")
            report_session.start_report_session()
            return {"status": "connected", "mode": "WIFI"}
            
        else:
            raise HTTPException(status_code=400, detail="Invalid DATA_SOURCE configured.")
            
    except Exception as e:
        logger.exception("Connect failed")
        raise HTTPException(status_code=500, detail=str(e))

@router.post("/disconnect")
def disconnect_system():
    if config.DATA_SOURCE == "SIMULATOR":
        if simulator.is_running():
            simulator.stop_simulator()
            sensors.stop_db_session()
            logger.info("Simulator stopped.")
    elif config.DATA_SOURCE == "SERIAL":
        sensors.stop_acquisition()
        sensors.stop_db_session()
        logger.info("Serial acquisition stopped.")
    elif config.DATA_SOURCE == "WIFI":
        sensors._running = False
        sensors.stop_db_session()
        logger.info("WiFi ingestion gate closed.")

    report_session.stop_report_session()
    return {"status": "disconnected"}

@router.post("/reset")
def reset_system():
    sensors.reset_system()
    return {"status": "system_reset"}

@router.post("/sensor/data/ingest")
def ingest_wifi_sensor_data(packet: SensorPacket):
    if not sensors.is_running():
        raise HTTPException(status_code=400, detail="System not running. Click Start System first.")
    
    # Convert Pydantic model back to dict for the pipeline
    data = packet.model_dump()
    
    # Debug logging as requested
    # logger.info(f"WiFi packet received: {data}")
    
    # Process the verified data
    sensors._process_data(data)
    return {"status": "ok"}

@router.post("/recording/start")
def start_recording():
    filename = sensors.start_recording()
    return {"status": "recording_started", "file": filename}

@router.post("/recording/stop")
def stop_recording():
    sensors.stop_recording()
    return {"status": "recording_stopped"}

@router.get("/state")
def get_full_state():
    return sensors.get_latest_state()

@router.get("/history")
def get_history_data():
    return {"history": sensors.get_history()}

@router.get("/map")
def get_slam_map():
    points = sensors.get_map_points()
    return {"points": [{"x": p[0], "y": p[1]} for p in points]}

@router.get("/alerts")
def get_alerts():
    return {"alerts": sensors.get_alerts()}

@router.get("/config/thresholds")
def get_thresholds():
    return {"thresholds": sensors.get_thresholds()}

@router.post("/config/thresholds")
def set_thresholds(thresholds: Dict[str, float]):
    sensors.set_thresholds(thresholds)
    return {"status": "updated", "thresholds": sensors.get_thresholds()}

# Sessions & History
@router.get("/sessions")
def list_sessions():
    """List all recorded inspection sessions."""
    return {"sessions": db_service.db.get_sessions()}

@router.get("/sessions/{session_id}")
def get_session(session_id: int):
    """Get metadata for a specific session."""
    summary = db_service.db.get_session_summary(session_id)
    if not summary:
        raise HTTPException(status_code=404, detail="Session not found")
    return summary

@router.delete("/sessions/{session_id}")
def delete_session(session_id: int):
    """Delete a session entirely from the database."""
    success = db_service.db.delete_session(session_id)
    if not success:
        raise HTTPException(status_code=500, detail="Failed to delete session")
    return {"status": "deleted", "session_id": session_id}

@router.get("/sessions/{session_id}/telemetry")
def get_session_telemetry(session_id: int):
    """Retrieve full telemetry data for a session (replay)."""
    data = db_service.db.get_session_telemetry(session_id)
    return {"telemetry": data}

@router.get("/sessions/{session_id}/report")
def get_session_report(session_id: int):
    """Generate and download a PDF inspection report (folder-based)."""
    try:
        pdf_path = report_service.report_gen.generate_folder_report()
        return FileResponse(
            path=pdf_path,
            filename=os.path.basename(pdf_path),
            media_type='application/pdf'
        )
    except Exception as e:
        logger.error(f"Error generating report: {e}")
        raise HTTPException(status_code=500, detail=str(e))

@router.api_route("/export-report", methods=["GET", "POST"])
def export_report():

    try:
        session_dir = report_session.get_report_dir()

        pdf_path = report_service.report_gen.generate_folder_report(session_dir)

        if not os.path.exists(pdf_path):
            raise Exception("PDF not generated")

        return FileResponse(
            path=pdf_path,
            media_type="application/pdf",
            filename="session_report.pdf"
        )

    except Exception as e:
        logger.error(f"Report generation failed: {e}")
        raise HTTPException(status_code=500, detail=str(e))

@router.get("/sessions/{session_id}/export/csv")
def export_session_csv(session_id: int):
    """Export session telemetry as a raw CSV file."""
    try:
        import pandas as pd
        telemetry = db_service.db.get_session_telemetry(session_id)
        if not telemetry:
            raise HTTPException(status_code=404, detail="No telemetry found for this session")
        
        df = pd.DataFrame(telemetry)
        export_dir = os.path.join(config.STORAGE_DIR, "Exports")
        os.makedirs(export_dir, exist_ok=True)
        
        csv_filename = f"Session_{session_id}_Telemetry.csv"
        csv_path = os.path.join(export_dir, csv_filename)
        df.to_csv(csv_path, index=False)
        
        return FileResponse(
            path=csv_path,
            filename=csv_filename,
            media_type='text/csv'
        )
    except Exception as e:
        logger.error(f"Error exporting CSV for session {session_id}: {e}")
        raise HTTPException(status_code=500, detail=str(e))

@router.post("/export/infringements")
def export_infringements():
    try:
        filename = sensors.export_infringement_history()
        return {"status": "exported", "file": filename}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@router.post("/export/acceleration")
def export_acceleration_data(req: AccelExportRequest):
    try:
        root = report_session.get_report_dir()
        filename = os.path.join(root, "acceleration.xlsx")
        
        from openpyxl import Workbook
        from openpyxl.styles import PatternFill
        
        wb = Workbook()
        ws = wb.active
        ws.title = "Acceleration Data"
        ws.append(["Timestamp", "Chainage", "Acc X", "Acc Y", "Acc Z", "Defect Status"])
        
        fills = {
            "UML": PatternFill(start_color="FFFF0000", end_color="FFFF0000", fill_type="solid"),
            "PML": PatternFill(start_color="FFFFFF00", end_color="FFFFFF00", fill_type="solid"),
            "CBML": PatternFill(start_color="FF00FF00", end_color="FF00FF00", fill_type="solid")
        }
        
        for pt in req.data:
            status = pt.status or "CBML"
            ws.append([pt.time, float(f"{pt.chainage:.3f}"), float(f"{pt.ax:.4f}"), float(f"{pt.ay:.4f}"), float(f"{pt.az:.4f}"), status])
            ws.cell(row=ws.max_row, column=6).fill = fills.get(status, fills["CBML"])
        
        wb.save(filename)
        return {"status": "exported", "file": filename}
    except Exception as e:
        logger.error(f"Export failed: {e}")
        raise HTTPException(status_code=500, detail=str(e))

# Video Feeds
@router.get("/video_feed")
def video_feed(index: int = 0):
    return StreamingResponse(cv_service.camera_manager.get_stream(index), media_type="multipart/x-mixed-replace; boundary=frame")

@router.get("/video_feed_geometry")
def video_feed_geometry(index: int = 3):
    # Left Frame: Raw Feed (as requested "output of webcam")
    return StreamingResponse(cv_service.camera_manager.get_stream(index), media_type="multipart/x-mixed-replace; boundary=frame")

@router.get("/video_feed_geometry_processed")
def video_feed_geometry_processed(index: int = 1):
    # Right Frame: Processed Script Output
    return StreamingResponse(cv_service.generate_track_geometry_feed(index), media_type="multipart/x-mixed-replace; boundary=frame")

@router.get("/video_feed_mask")
def video_feed_mask(index: int = 1):
    return StreamingResponse(cv_service.generate_mask_feed(index), media_type="multipart/x-mixed-replace; boundary=frame")

@router.get("/video_feed_yolo")
def video_feed_yolo(index: int = 2):
    return StreamingResponse(cv_service.generate_yolo_feed(index), media_type="multipart/x-mixed-replace; boundary=frame")

@router.get("/video_feed_rail_ai")
def video_feed_rail_ai(index: int = 2):
    return StreamingResponse(cv_service.generate_rail_ai_feed(index), media_type="multipart/x-mixed-replace; boundary=frame")

@router.get("/video_feed_condition_overlay")
def video_feed_condition_overlay(index: int = 1):
    return StreamingResponse(cv_service.generate_condition_feed_overlay(index), media_type="multipart/x-mixed-replace; boundary=frame")

@router.get("/video_feed_condition_mask")
def video_feed_condition_mask(index: int = 1):
    return StreamingResponse(cv_service.generate_condition_feed_mask(index), media_type="multipart/x-mixed-replace; boundary=frame")

@router.get("/geometry/data")
def get_geometry_data():
    return {
        "timestamp": time.time(),
        "chainage": get_current_chainage(),
        "gauge_pixels": cv_service.geo_processor.latest_gauge,
        "running": sensors.is_running(),
        "status": "active" if sensors.is_running() else "idle"
    }

@router.post("/recording/geometry/start")
def start_geometry_recording():
    cv_service.geo_processor.start_recording()
    return {"status": "started"}

@router.post("/recording/geometry/stop")
def stop_geometry_recording():
    cv_service.geo_processor.stop_recording()
    return {"status": "stopped"}

@router.post("/recording/condition/start")
def start_condition_recording():
    cv_service.yolo_processor.start_recording()
    return {"status": "started"}

@router.post("/recording/condition/stop")
def stop_condition_recording():
    cv_service.yolo_processor.stop_recording()
    return {"status": "stopped"}

@router.post("/recording/rearwindow/start")
def start_rearwindow_recording():
    cv_service.rear_recorder.start_recording()
    return {"status": "started"}

@router.post("/recording/rearwindow/stop")
def stop_rearwindow_recording():
    cv_service.rear_recorder.stop_recording()
    return {"status": "stopped"}

@router.get("/condition/mode")
def get_condition_mode():
    return {"mode": CONDITION_MODE}

@router.post("/condition/mode")
def set_condition_mode(update: ModeUpdate):
    global CONDITION_MODE
    if update.mode not in ["laser", "yolo"]:
        raise HTTPException(status_code=400, detail="Invalid mode")
    CONDITION_MODE = update.mode
    return {"status": "updated", "mode": CONDITION_MODE}

@router.post("/condition/threshold")
def set_condition_threshold(update: ThresholdUpdate):
    cv_service.cond_processor.threshold = update.threshold
    return {"status": "updated", "threshold": cv_service.cond_processor.threshold}

@router.get("/models/list")
def list_models():
    search_paths = [
        os.path.join(os.getcwd(), "models"),
        os.path.join(os.getcwd(), "runs"),
        os.getcwd()
    ]
    candidates = []
    for root_path in search_paths:
        if os.path.exists(root_path):
            for root, dirs, files in os.walk(root_path):
                if "best.pt" in files:
                    candidates.append(os.path.join(root, "best.pt"))
                if "models" in root_path:
                    for f in files:
                        if f.endswith(".pt") and f != "best.pt":
                            candidates.append(os.path.join(root, f))
    return {"models": list(set(candidates))}

@router.get("/models_rail/list")
def list_models_rail():
    rail_models_path = os.path.join(os.getcwd(), "models_rail")
    candidates = []
    if os.path.exists(rail_models_path):
        for root, dirs, files in os.walk(rail_models_path):
            for f in files:
                if f.endswith(".pt"):
                    candidates.append(os.path.join(root, f))
    return {"models": candidates}

@router.post("/models/select")
def select_model(selection: ModelSelect):
    success = cv_service.yolo_processor.load_model(selection.path)
    if success:
        return {"status": "loaded", "path": selection.path}
    else:
        raise HTTPException(status_code=500, detail="Failed to load model")

@router.get("/camera/list")
def list_cameras():
    cameras = cv_service.get_available_cameras()
    return {"cameras": cameras}

@router.post("/reset")
def reset_system():
    """Reset the distance/tracker state in the backend."""
    sensors.reset_system()
    return {"status": "success", "message": "System reset"}

@router.post("/camera/start")
def start_camera(req: CameraControl):
    # Camera generation handles lifecycle automatically
    return {"status": "camera_ready", "index": req.index}

@router.post("/camera/stop")
def stop_camera(req: CameraControl):
    cv_service.camera_manager.force_stop(req.index)
    return {"status": "camera_stopped", "index": req.index}

@router.get("/video_feed_face")
def video_feed_face(index: int = 0):
    return StreamingResponse(cv_service.camera_manager.get_stream(index), media_type="multipart/x-mixed-replace; boundary=frame")

# --------------------------------------------------------------------------
# SIMULATOR CONTROL
# --------------------------------------------------------------------------

@router.post("/sim/start")
def start_sim(rate_hz: int = Body(20, embed=True)):
    """Start the sensor simulator at the specified Hz (default 20)."""
    if simulator.is_running():
        return {"status": "already_running"}
    simulator.start_simulator(rate_hz=rate_hz)
    return {"status": "simulator_started", "rate_hz": rate_hz}

@router.post("/sim/stop")
def stop_sim():
    """Stop the sensor simulator."""
    if not simulator.is_running():
        return {"status": "not_running"}
    simulator.stop_simulator()
    return {"status": "simulator_stopped"}

@router.websocket("/ws")
async def websocket_handler(websocket: WebSocket):
    await sensor_service.manager.connect(websocket)
    try:
        while True:
            data = await websocket.receive_text()
            if data:
                try:
                    msg = json.loads(data)
                    if msg.get("cmd") == "reset":
                        sensors.reset_system()
                        await websocket.send_json({"type": "info", "msg": "Reset performed"})
                except: pass
    except WebSocketDisconnect:
        sensor_service.manager.disconnect(websocket)
